#!/usr/bin/env python3
"""
compare_migration.py -- SysML v1.7 -> SysML v2 migration success auditor for the CATIA Magic / MSOSA 2026x SysML v2 migration plugin.

Part of: Metamodel Scope Using Model-Based Decision on Intended-Use & Requirements Engineering Categorization Tool (MB-DIRECT), Doctor of Engineering praxis, The George Washington
         University, School of Engineering and Applied Science, 2026.
Author:  Evan Matthews (ORCID 0009-0005-2265-0754)
License: MIT (see LICENSE in the repository root)
AI assistance: developed with assistance from Anthropic's Claude (Claude Code);
         specified, reviewed, tested, and validated by the author in
         accordance with the AI Policy for George Washington University Online Engineering Programs.

This single file consolidates the three modules of the praxis Migration
compare_migration.py : the base comparison, plus a per-metatype breakdown on the summary sheet, with a
fixed set of metatypes excluded from every metric.

Workflow
--------
Step 1 (build):   python compare_migration.py AllElements_1_7.csv
    Takes the CSV produced by the AllElements_Export.vm Report Wizard
    template (or the 33-column TSV exported from MSOSA described in tsv_to_sysmlv2.py) and builds
    <name>_MigrationComparison.xlsx containing:
      - CONFIG              (two yellow cells: paths to the CATIA
                             MigrationReportOf<project>.xlsx and to the
                             exported SysML v2 .txt/.sysml source)
      - ALL_ELEMENTS_1_7    (the v1.7 element/diagram inventory, minus
                             the excluded metatypes)
      - EXCLUDED_METATYPES  (the rows set aside, same columns)
    If a MigrationReportOf*.xlsx or a single .txt/.sysml file sits in
    the same folder as the CSV, the paths are pre-filled automatically
    and the comparison runs immediately.

Step 2 (compare): python compare_migration.py <workbook>_MigrationComparison.xlsx
    Reads the two file paths from CONFIG, then creates/overwrites:
      - MIGRATED_SUCCESS   every v1.7 element found in the SysML v2
                           source, with the v2 construct it became,
                           its v2 qualified name, line number, match
                           confidence, and any properties the CATIA
                           report says were lost.
      - MISSED_MIGRATIONS  elements with no trace in the v2 source that
                           are ALSO absent from the CATIA not-migrated
                           report (categories: MISSING, REFERENCE_ONLY,
                           UNVERIFIABLE_UNNAMED).
      - MIGRATION_SUMMARY  counts and coverage percentages, a note on
                           how many elements were excluded, and a
                           per-metatype section (Element Metatype |
                           Quantity | % Migration Success) driven by
                           live COUNTIF formulas.

Excluded metatypes (see EXCLUDED_METATYPES below): Boolean Tagged Value,
Diagram, Element Tagged Value, Integer Tagged Value, Literal Unlimited
Natural, String Tagged Value. They are moved out of ALL_ELEMENTS_1_7
before the comparison runs, so every downstream number is computed
without them. Re-running on the same workbook is safe: rows already on
the EXCLUDED_METATYPES sheet are kept, deduplicated by Element ID.

Optional flags (both modes):
    --catia <path>   path to MigrationReportOf<project>.xlsx (also written to CONFIG)
    --v2 <path>      path to the SysML v2 .txt/.sysml source (also written to CONFIG)

Dependencies: Python 3.8+ and openpyxl (pip install openpyxl).
"""

import argparse
import csv
import datetime
import glob
import io
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- constants

SHEET_CONFIG = "CONFIG"
SHEET_ALL = "ALL_ELEMENTS_1_7"
SHEET_SUCCESS = "MIGRATED_SUCCESS"
SHEET_MISSED = "MISSED_MIGRATIONS"
SHEET_SUMMARY = "MIGRATION_SUMMARY"

CONFIG_LABEL_CATIA = "CATIA migration report (.xlsx):"
CONFIG_LABEL_V2 = "SysML v2 source code (.txt/.sysml):"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
CAT_FILLS = {
    "MISSING": PatternFill("solid", fgColor="FFC7CE"),
    "REFERENCE_ONLY": PatternFill("solid", fgColor="FFEB9C"),
    "UNVERIFIABLE_UNNAMED": PatternFill("solid", fgColor="DDEBF7"),
}

ALL_HEADERS = ["Element MetaType", "Element Name", "Qualified Name", "Element ID",
               "Kind", "Owner Qualified Name", "Applied Stereotypes"]

# SysML v2 declaration keywords that can introduce a (possibly named) element.
# Two-token keywords are matched before single-token ones; a following 'def'
# is always absorbed (e.g. "part def", "use case def", "attribute def").
KW_TWO = {("use", "case")}
KW_ONE = {
    "package", "part", "attribute", "port", "connection", "interface",
    "action", "state", "transition", "requirement", "constraint", "item",
    "comment", "doc", "occurrence", "event", "flow", "message", "succession",
    "binding", "view", "viewpoint", "concern", "calc", "enum", "analysis",
    "verification", "individual", "metadata", "rendering", "allocation",
    "objective", "subject", "actor", "stakeholder", "satisfy", "verify",
    "allocate", "perform", "exhibit", "include", "assert", "assume",
    "require", "frame", "snapshot", "timeslice", "case", "accept", "send",
    "dependency", "def",
    "merge", "decide", "fork", "join",   # control nodes (MSOSA emits them named)
}
# Modifiers that may precede the declaration keyword.
MODIFIERS = {
    "public", "private", "protected", "abstract", "variation", "ref",
    "readonly", "derived", "end", "in", "out", "inout", "nonunique",
    "ordered", "entry", "exit", "do", "first", "all", "individual",
}
# Tokens that can follow a keyword but are NOT a declared name.
NOT_A_NAME = {
    "connect", "specializes", "redefines", "subsets", "references", "of",
    "to", "from", "then", "first", "at", "about", "by", "as", "defined",
    "ref", "all", "def", "if", "when", "via", "assign", "new", "istype",
    "hastype", "meta", "locale", "bind", "until", "after",
} | KW_ONE

TOKEN_RE = re.compile(r"'(?:[^'\\]|\\.)*'|[A-Za-z_][A-Za-z0-9_]*|\S")
IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
MDEL_ID_RE = re.compile(r"mdel://([^?]+)")

V2_ESCAPES = {"n": "\n", "t": "\t", "r": "\r", "b": "\b", "f": "\f",
              "'": "'", '"': '"', "\\": "\\"}


# ------------------------------------------------------------ normalization

def canonical_name(s):
    """Normalize an element name so v1.7, CATIA-report and SysML v2
    spellings of the same name compare equal."""
    if s is None:
        return ""
    s = str(s)
    s = s.replace("<NL>", "\n")            # velocity template line-break token
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    return s.strip()


def unquote_v2(tok):
    """Turn a SysML v2 name token (possibly 'quoted' with escapes) into a
    plain string."""
    if len(tok) >= 2 and tok.startswith("'") and tok.endswith("'"):
        body = tok[1:-1]
        out, i = [], 0
        while i < len(body):
            c = body[i]
            if c == "\\" and i + 1 < len(body):
                out.append(V2_ESCAPES.get(body[i + 1], body[i + 1]))
                i += 2
            else:
                out.append(c)
                i += 1
        return "".join(out)
    return tok


def canonical_qn(qn):
    """Normalize a qualified name: split on '::', drop empty segments
    (unnamed owners), canonicalize each segment, rejoin."""
    if qn is None:
        return ""
    parts = [canonical_name(p) for p in str(qn).replace("<NL>", "\n").split("::")]
    return "::".join(p for p in parts if p)


# ------------------------------------------------------- SysML v2 text parse

def strip_comments(text):
    """Replace // line comments and /* */ block comments (incl. doc bodies)
    with spaces, preserving newlines and quoted names."""
    out = list(text)
    i, n = 0, len(text)
    state = "code"
    while i < n:
        c = text[i]
        nxt = text[i + 1] if i + 1 < n else ""
        if state == "code":
            if c == "'":
                state = "quote"
            elif c == "/" and nxt == "/":
                state = "line"
                out[i] = out[i + 1] = " "
                i += 1
            elif c == "/" and nxt == "*":
                state = "block"
                out[i] = out[i + 1] = " "
                i += 1
        elif state == "quote":
            if c == "\\":
                i += 1
            elif c == "'":
                state = "code"
        elif state == "line":
            if c == "\n":
                state = "code"
            else:
                out[i] = " "
        elif state == "block":
            if c == "*" and nxt == "/":
                # close the comment AND terminate the statement: doc/comment
                # bodies (`doc /* ... */`) have no ';' of their own, and
                # without this the next declaration would merge into them
                out[i] = " "
                out[i + 1] = ";"
                state = "code"
                i += 1
            elif c != "\n":
                out[i] = " "
        i += 1
    return "".join(out)


def parse_statement(tokens):
    """Given the tokens of one statement, return (construct, name) where
    name may be '' for unnamed declarations, or (None, None) if the
    statement declares nothing we track."""
    i = 0
    mods = []
    # metadata-prefixed declarations, e.g. MSOSA 2026x's
    #   #RequirementDerivation::derivation connection '54' { ... }
    # skip the '#'/'@' marker and the metadata qualified name after it, then
    # parse the remainder as a normal statement (stops at a real keyword)
    while i < len(tokens) and tokens[i] in ("#", "@"):
        i += 1
        while i < len(tokens) and (
                tokens[i] == ":"
                or (tokens[i] not in KW_ONE and tokens[i] not in MODIFIERS
                    and (tokens[i].startswith("'")
                         or IDENT_RE.match(tokens[i])))):
            i += 1
    tokens = tokens[i:]
    i = 0
    # skip leading modifiers
    while i < len(tokens) and tokens[i] in MODIFIERS:
        # 'end'/'ref' etc. can themselves head a usage; if the *next* token
        # is not a keyword/modifier, treat the current token as the keyword.
        if (tokens[i] in ("end", "ref") and i + 1 < len(tokens)
                and tokens[i + 1] not in MODIFIERS
                and tokens[i + 1] not in KW_ONE
                and (tokens[i + 1],) not in [(a,) for a, b in KW_TWO]):
            break
        mods.append(tokens[i])
        i += 1
    if i >= len(tokens):
        return None, None
    # keyword CHAIN (e.g. "assert constraint", "satisfy requirement",
    # "perform action", "event occurrence", "use case def"), absorbing 'def'
    parts = []
    while i < len(tokens) and len(parts) < 3:
        if i + 1 < len(tokens) and (tokens[i], tokens[i + 1]) in KW_TWO:
            parts += [tokens[i], tokens[i + 1]]
            i += 2
        elif tokens[i] in KW_ONE and tokens[i] != "def":
            parts.append(tokens[i])
            i += 1
        elif tokens[i] in ("end", "ref") and not parts:
            parts.append(tokens[i])
            i += 1
        else:
            break
        if i < len(tokens) and tokens[i] == "def":
            parts.append("def")
            i += 1
            break
    if parts:
        construct = " ".join(parts)
    elif mods and set(mods) & {"in", "out", "inout"} and i < len(tokens) and (
            tokens[i].startswith("'") or
            (IDENT_RE.match(tokens[i]) and tokens[i] not in NOT_A_NAME)):
        # keyword-less parameter declaration, e.g.  "in accelPosition;"
        construct = [m for m in mods if m in ("in", "out", "inout")][-1] + " parameter"
    else:
        return None, None
    # optional short name  <shortName>
    if i < len(tokens) and tokens[i] == "<":
        while i < len(tokens) and tokens[i] != ">":
            i += 1
        i += 1
    # optional multiplicity  [..]
    if i < len(tokens) and tokens[i] == "[":
        while i < len(tokens) and tokens[i] != "]":
            i += 1
        i += 1
    # declared name?
    name = ""
    if i < len(tokens):
        tok = tokens[i]
        if tok.startswith("'"):
            name = unquote_v2(tok)
        elif IDENT_RE.match(tok) and tok not in NOT_A_NAME and tok not in MODIFIERS:
            name = tok
    return construct, canonical_name(name)


class V2Model:
    def __init__(self):
        self.by_qn = {}      # canonical qn -> list of (construct, line, qn)
        self.by_name = {}    # canonical name -> list of (construct, line, qn)
        self.tokens = set()  # every identifier / quoted-name in the file
        self.unnamed = {}    # construct -> count of unnamed declarations
        self.total_decls = 0


def parse_sysml_v2(text):
    """Single pass over the v2 source: track brace nesting to build
    qualified names, record every (possibly named) declaration."""
    cleaned = strip_comments(text)
    model = V2Model()
    for tok in TOKEN_RE.findall(cleaned):
        if tok.startswith("'"):
            model.tokens.add(canonical_name(unquote_v2(tok)))
        elif IDENT_RE.match(tok):
            model.tokens.add(tok)

    stack = []          # names of enclosing scopes ('' for unnamed)
    cstack = []         # constructs of enclosing scopes (e.g. 'enum def')
    buf = []            # current statement characters
    buf_line = 1        # line where current statement started
    line = 1
    in_quote = False
    i, n = 0, len(cleaned)

    def flush(opens_scope):
        nonlocal buf, buf_line
        stmt = "".join(buf).strip()
        buf = []
        name, construct = "", None
        if stmt:
            tokens = TOKEN_RE.findall(stmt)
            construct, name = parse_statement(tokens)
            # bare-name statements inside an enum def are enumeration literals
            if (construct is None and not opens_scope and cstack
                    and (cstack[-1] or "").startswith("enum") and tokens
                    and (tokens[0].startswith("'") or
                         (IDENT_RE.match(tokens[0]) and tokens[0] not in NOT_A_NAME))
                    and (len(tokens) == 1 or tokens[1] == "=")):
                construct = "enum literal"
                name = canonical_name(unquote_v2(tokens[0]))
            if construct:
                model.total_decls += 1
                qn_parts = [p for p in stack if p]
                if name:
                    qn = "::".join(qn_parts + [name])
                    rec = (construct, buf_line, qn)
                    model.by_qn.setdefault(qn, []).append(rec)
                    model.by_name.setdefault(name, []).append(rec)
                else:
                    model.unnamed[construct] = model.unnamed.get(construct, 0) + 1
                # an accept clause's payload is a real nested feature of the
                # action ("action '57' accept result : Sig;" declares
                # '57'::result) - record it so 1.7 result pins can match
                if name and not construct.startswith("accept"):
                    for j, t in enumerate(tokens[:-1]):
                        if t != "accept":
                            continue
                        nxt = tokens[j + 1]
                        pname = ""
                        if nxt.startswith("'"):
                            pname = canonical_name(unquote_v2(nxt))
                        elif (IDENT_RE.match(nxt) and nxt not in NOT_A_NAME
                              and nxt not in MODIFIERS):
                            pname = canonical_name(nxt)
                        if pname:
                            pqn = "::".join(qn_parts + [name, pname])
                            prec = ("accept payload", buf_line, pqn)
                            model.by_qn.setdefault(pqn, []).append(prec)
                            model.by_name.setdefault(pname, []).append(prec)
                        break
        if opens_scope:
            stack.append(name)
            cstack.append(construct)
        buf_line = line

    while i < n:
        c = cleaned[i]
        if c == "\n":
            line += 1
        if in_quote:
            buf.append(c)
            if c == "\\":
                if i + 1 < n:
                    buf.append(cleaned[i + 1])
                    i += 1
            elif c == "'":
                in_quote = False
        elif c == "'":
            in_quote = True
            buf.append(c)
        elif c == "{":
            flush(opens_scope=True)
        elif c == "}":
            flush(opens_scope=False)
            if stack:
                stack.pop()
            if cstack:
                cstack.pop()
        elif c == ";":
            flush(opens_scope=False)
        else:
            buf.append(c)
        i += 1
    flush(opens_scope=False)
    return model


# ---------------------------------------------------- CATIA report parsing

def _header_map(ws):
    return {str(c.value).strip(): idx
            for idx, c in enumerate(ws[1]) if c.value is not None}


def parse_catia_report(path):
    """Return (elements, props):
    elements: element id -> {metatype, name, qn}
    props:    element id -> [not-migrated property names]"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    elements, props = {}, {}

    el_sheet = next((s for s in wb.sheetnames
                     if "not migrated element" in s.lower()
                     and "propert" not in s.lower()), None)
    pr_sheet = next((s for s in wb.sheetnames
                     if "propert" in s.lower()), None)
    if el_sheet is None:
        raise ValueError(
            f"'{os.path.basename(path)}' has no 'Not Migrated Elements' sheet "
            f"(sheets found: {wb.sheetnames}). Is this really the CATIA "
            "migration report?")

    def ingest_element_sheet(ws, hm, default_metatype=None):
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            eid = row[hm.get("Element ID", 4)] if len(row) > hm.get("Element ID", 4) else None
            if not eid:
                url = row[hm.get("Element Url", 3)] if "Element Url" in hm else None
                m = MDEL_ID_RE.search(str(url or ""))
                eid = m.group(1) if m else None
            if not eid:
                continue
            elements[str(eid)] = {
                "metatype": default_metatype
                            or row[hm.get("Element MetaType", 0)] or "",
                "name": canonical_name(row[hm.get("Element Name", 1)]),
                "qn": canonical_qn(row[hm.get("Qualified Name", 2)]),
            }

    ws = wb[el_sheet]
    ingest_element_sheet(ws, _header_map(ws))

    # tsv_to_sysmlv2.py reports its unmigrated diagrams on a separate
    # 'Not Migrated Diagrams' sheet (the CATIA plugin lists them inline in
    # the elements sheet) - ingest it too so both report styles produce the
    # same 'excluded' accounting for diagrams
    dg_sheet = next((s for s in wb.sheetnames
                     if "not migrated diagram" in s.lower()), None)
    if dg_sheet:
        ws = wb[dg_sheet]
        ingest_element_sheet(ws, _header_map(ws), default_metatype="Diagram")

    if pr_sheet:
        ws = wb[pr_sheet]
        hm = _header_map(ws)
        url_col = hm.get("Element Url")
        prop_col = hm.get("Property Name")
        cur_id = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            url = row[url_col] if url_col is not None and len(row) > url_col else None
            if url:  # a new element starts; blank url rows continue the previous one
                m = MDEL_ID_RE.search(str(url))
                cur_id = m.group(1) if m else None
            prop = row[prop_col] if prop_col is not None and len(row) > prop_col else None
            if cur_id and prop:
                props.setdefault(cur_id, []).append(str(prop))
    wb.close()
    return elements, props


# ------------------------------------------------- velocity CSV / ALL sheet

def _scan_quoted_records(s, ncols):
    """Structural parser for run-together CSV rows (produced by the pre-fix
    template, where Report Wizard swallowed the newline after each row).
    Every field is quoted; fields within a record are comma-separated; a
    record boundary is a closing quote immediately followed by an opening
    quote. '""' inside fields 1..ncols-1 is an escaped quote; in the last
    field (Applied Stereotypes, which never contains quotes) it closes the
    record."""
    recs, cur, i, n = [], [], 0, len(s)
    while i < n:
        c = s[i]
        if c != '"':
            i += 1  # separators (commas/newlines) and stray bytes
            continue
        i += 1
        field = []
        while i < n:
            c = s[i]
            if c == '"':
                if i + 1 < n and s[i + 1] == '"' and len(cur) < ncols - 1:
                    field.append('"')
                    i += 2
                    continue
                i += 1
                break
            field.append(c)
            i += 1
        cur.append("".join(field))
        if len(cur) == ncols:
            recs.append(cur)
            cur = []
    if cur:
        recs.append(cur + [""] * (ncols - len(cur)))
    return recs


def _read_migration_tsv(text, path):
    """Parse the 33-column SysML17_Migration_Export[_w_Diagrams].vm TSV and
    map it onto the standard element-dict shape. Tolerates the pre-v6
    template's blank-line flood and drops fragment rows produced by
    unsanitized line breaks in names."""
    rows = [ln.split("\t") for ln in text.splitlines() if ln.strip()]
    hdr = next((r for r in rows if "SysML 1.7 Element Type" in r), None)
    if hdr is None:
        raise ValueError(f"{path}: could not find the migration-export header row")
    idx = {h.strip(): i for i, h in enumerate(hdr)}

    def get(r, col):
        j = idx.get(col)
        return r[j].strip() if j is not None and j < len(r) else ""

    elements, dropped = [], 0
    for r in rows[rows.index(hdr) + 1:]:
        if len(r) < 6:
            dropped += 1
            continue
        mt = get(r, "SysML 1.7 Element Type")
        qn = get(r, "Qualified Name")
        if not mt and not qn:
            dropped += 1
            continue
        elements.append({
            "metatype": mt,
            "name": get(r, "Name"),
            "qn": qn,
            "id": get(r, "Element ID"),
            "kind": "Diagram" if mt == "Diagram" else "Element",
            "owner": get(r, "Owner Qualified Name"),
            "stereotypes": get(r, "Applied Stereotype"),
        })
    print(f"Read {len(elements)} elements from migration-export TSV "
          f"{os.path.basename(path)}" +
          (f" (skipped {dropped} blank/fragment rows)" if dropped else ""))
    if elements and not any(e["id"] for e in elements):
        print("NOTE: 'Element ID' column is empty (pre-v6 template) - CATIA "
              "report matching will fall back to qualified names. Regenerate "
              "with the v6 SysML17_Migration_Export_w_Diagrams.vm for exact "
              "ID matching.")
    return elements


def read_velocity_csv(path):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with open(path, encoding=enc) as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Could not decode {path} as UTF-8 or cp1252")

    if "SysML 1.7 Element Type" in text[:20000]:
        return _read_migration_tsv(text, path)

    physical_lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(physical_lines) <= 3 and len(text) > 5000:
        # old-template output: rows ran together on one line - recover
        header = next(csv.reader([physical_lines[0][:20000]]))
        header = [h.strip() for h in header[:len(ALL_HEADERS)]]
        body = text[text.find("\n") + 1:]
        data = _scan_quoted_records(body, len(header))
        data = [r for r in data if r[:3] != header[:3]]  # drop duplicated header
        rows = [header] + data
        print(f"NOTE: {os.path.basename(path)} had all rows on one line "
              f"(old template); structurally recovered {len(data)} records. "
              "Regenerate with the current AllElements_Export.vm to avoid this.")
    else:
        rows = list(csv.reader(io.StringIO(text)))
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError(f"{path} is empty")
    header = [h.strip() for h in rows[0]]
    if "Qualified Name" not in header or "Element MetaType" not in header:
        raise ValueError(
            f"{path} does not look like an AllElements_Export.vm output "
            f"(header row found: {header})")
    idx = {h: header.index(h) for h in header}

    def get(r, col):
        j = idx.get(col)
        return r[j] if j is not None and j < len(r) else ""

    elements = []
    for r in rows[1:]:
        elements.append({
            "metatype": get(r, "Element MetaType"),
            "name": get(r, "Element Name"),
            "qn": get(r, "Qualified Name"),
            "id": get(r, "Element ID"),
            "kind": get(r, "Kind") or "Element",
            "owner": get(r, "Owner Qualified Name"),
            "stereotypes": get(r, "Applied Stereotypes"),
        })
    return elements


def read_all_sheet(wb):
    if SHEET_ALL not in wb.sheetnames:
        raise ValueError(f"Workbook has no {SHEET_ALL} sheet. Run this script "
                         "on the velocity CSV first to build the workbook.")
    ws = wb[SHEET_ALL]
    hm = _header_map(ws)
    elements = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        def get(col, default=""):
            j = hm.get(col)
            v = row[j] if j is not None and j < len(row) else None
            return "" if v is None else str(v)
        elements.append({
            "metatype": get("Element MetaType"),
            "name": get("Element Name"),
            "qn": get("Qualified Name"),
            "id": get("Element ID"),
            "kind": get("Kind") or "Element",
            "owner": get("Owner Qualified Name"),
            "stereotypes": get("Applied Stereotypes"),
        })
    return elements


# ------------------------------------------------------------- xlsx styling

def style_header(ws, ncols):
    for j in range(1, ncols + 1):
        c = ws.cell(row=1, column=j)
        c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(ws.max_row, 1)}"


def autosize(ws, ncols, cap=70):
    for j in range(1, ncols + 1):
        width = 10
        for cell in ws[get_column_letter(j)]:
            v = cell.value
            if v is not None:
                width = max(width, min(cap, len(str(v)) + 2))
        ws.column_dimensions[get_column_letter(j)].width = width


def write_table(ws, headers, rows):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT_NAME, size=10)
    style_header(ws, len(headers))
    autosize(ws, len(headers))


# ------------------------------------------------------------ build (step 1)

def find_config_row(ws, label):
    for row in ws.iter_rows(min_col=1, max_col=1):
        cell = row[0]
        if cell.value and str(cell.value).strip().lower().startswith(
                label.split("(")[0].strip().lower()):
            return cell.row
    return None


def build_workbook(csv_path, catia=None, v2=None):
    elements = read_velocity_csv(csv_path)
    folder = os.path.dirname(os.path.abspath(csv_path))

    if not catia:
        cands = glob.glob(os.path.join(folder, "MigrationReportOf*.xlsx"))
        catia = cands[0] if len(cands) == 1 else ""
    if not v2:
        cands = [p for p in glob.glob(os.path.join(folder, "*.txt"))
                 + glob.glob(os.path.join(folder, "*.sysml"))]
        v2 = cands[0] if len(cands) == 1 else ""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_CONFIG
    ws["A1"] = ("CONFIG - fill in the two yellow cells below, save, then run:  "
                "python compare_migration.py <this workbook>")
    ws["A1"].font = Font(name=FONT_NAME, size=11, bold=True)
    ws["A3"] = CONFIG_LABEL_CATIA
    ws["A4"] = CONFIG_LABEL_V2
    ws["B3"] = catia
    ws["B4"] = v2
    for addr in ("A3", "A4"):
        ws[addr].font = Font(name=FONT_NAME, size=10, bold=True)
    for addr in ("B3", "B4"):
        ws[addr].fill = INPUT_FILL
        ws[addr].font = Font(name=FONT_NAME, size=10)
    ws["A6"] = "Legend: yellow cells are the only cells you need to edit (full file paths)."
    ws["A7"] = (r"Example:  C:\models\MigrationReportOfmyproject.xlsx   /   "
                r"C:\models\package MyModel.txt")
    for addr in ("A6", "A7"):
        ws[addr].font = Font(name=FONT_NAME, size=9, italic=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 110

    ws2 = wb.create_sheet(SHEET_ALL)
    write_table(ws2, ALL_HEADERS,
                [[e["metatype"], e["name"], e["qn"], e["id"], e["kind"],
                  e["owner"], e["stereotypes"]] for e in elements])

    out = os.path.splitext(os.path.abspath(csv_path))[0] + "_MigrationComparison.xlsx"
    wb.save(out)
    print(f"Built workbook: {out}")
    print(f"  {len(elements)} v1.7 elements/diagrams imported into {SHEET_ALL}.")
    print(f"  CONFIG paths pre-filled: CATIA report = {catia or '(fill in)'}")
    print(f"                           SysML v2 src = {v2 or '(fill in)'}")
    print("Next: check/fill CONFIG!B3 and CONFIG!B4, then run:")
    print(f'  python compare_migration.py "{out}"')
    return out


# ---------------------------------------------------------- compare (step 2)

def classify(elements, catia_elements, catia_props, v2m):
    """Return (success_rows, missed_rows, stats)."""
    catia_by_qn = {}
    for eid, rec in catia_elements.items():
        key = (rec["qn"], rec["name"])
        catia_by_qn.setdefault(key, set()).add(eid)

    success, missed = [], []
    stats = {"total": 0, "diagrams": 0, "in_catia": 0, "matched_qn": 0,
             "matched_name": 0, "reference_only": 0, "missing": 0,
             "unnamed": 0, "also_in_catia_but_found": 0}

    for e in elements:
        stats["total"] += 1
        name = canonical_name(e["name"])
        qn = canonical_qn(e["qn"])
        eid = (e["id"] or "").strip()
        if e["kind"] == "Diagram":
            stats["diagrams"] += 1

        in_catia = bool(eid and eid in catia_elements)
        if not in_catia and (qn, name) in catia_by_qn and not eid:
            in_catia = True  # fallback when the export carries no IDs

        # --- unnamed elements cannot be traced by name/QN -------------------
        # (their qualified name collapses to the owner's, which would
        #  otherwise false-match the owner's v2 declaration)
        if not name:
            if in_catia:
                stats["in_catia"] += 1
            else:
                stats["unnamed"] += 1
                missed.append(["UNVERIFIABLE_UNNAMED", e["metatype"], e["name"],
                               e["qn"], eid,
                               "Unnamed in v1.7 - cannot be verified by name "
                               "matching; review manually"])
            continue

        # --- diagrams can never live in textual v2: name-matching them only
        # ever hits a same-named model element (e.g. the diagram named after
        # its owning activity matching that activity's def) -----------------
        if e["kind"] == "Diagram" or e["metatype"] == "Diagram":
            if in_catia:
                stats["in_catia"] += 1
            else:
                stats["missing"] += 1
                missed.append(["MISSING", e["metatype"], e["name"], e["qn"],
                               eid,
                               "Diagram - SysML v2 text carries no diagrams, "
                               "and it is not listed in the migration "
                               "report's not-migrated sheets either"])
            continue

        # --- look for the element in the SysML v2 source -------------------
        match, match_type = None, ""
        if qn and qn in v2m.by_qn:
            match, match_type = v2m.by_qn[qn], "Qualified name"
        elif name and name in v2m.by_name:
            recs = v2m.by_name[name]
            match = recs
            match_type = ("Name (unique)" if len({r[2] for r in recs}) == 1
                          else "Name (ambiguous)")

        lost = catia_props.get(eid, [])

        if match:
            constructs = "; ".join(sorted({r[0] for r in match}))
            v2qns = "; ".join(sorted({r[2] for r in match}))
            lines = "; ".join(str(r[1]) for r in match[:8])
            if match_type == "Qualified name":
                stats["matched_qn"] += 1
            else:
                stats["matched_name"] += 1
            notes = []
            if in_catia:
                notes.append("Also listed in CATIA not-migrated report - verify manually")
                stats["also_in_catia_but_found"] += 1
            if lost:
                notes.append("Properties not migrated: " + ", ".join(sorted(set(lost))))
            success.append([e["metatype"], e["name"], e["qn"], eid,
                            constructs, v2qns, lines, match_type,
                            "; ".join(notes)])
            continue

        if in_catia:
            stats["in_catia"] += 1
            continue  # accounted for by the CATIA report itself

        # --- not in v2, not in the CATIA report -> MISSED ------------------
        if name in v2m.tokens:
            stats["reference_only"] += 1
            missed.append(["REFERENCE_ONLY", e["metatype"], e["name"],
                           e["qn"], eid,
                           "Name appears in the v2 source only as a "
                           "reference, not as a declaration - review manually"])
        else:
            stats["missing"] += 1
            missed.append(["MISSING", e["metatype"], e["name"], e["qn"], eid,
                           "No trace in the SysML v2 source and absent from "
                           "the CATIA not-migrated report"])

    order = {"MISSING": 0, "REFERENCE_ONLY": 1, "UNVERIFIABLE_UNNAMED": 2}
    missed.sort(key=lambda r: (order.get(r[0], 9), r[3], r[2]))
    success.sort(key=lambda r: (r[2], r[1]))
    return success, missed, stats


def run_comparison(xlsx_path, catia=None, v2=None):
    wb = openpyxl.load_workbook(xlsx_path)
    if SHEET_CONFIG not in wb.sheetnames:
        raise ValueError(f"Workbook has no {SHEET_CONFIG} sheet. Run the script "
                         "on the velocity CSV first to build the workbook.")
    cfg = wb[SHEET_CONFIG]
    r_catia = find_config_row(cfg, CONFIG_LABEL_CATIA) or 3
    r_v2 = find_config_row(cfg, CONFIG_LABEL_V2) or 4
    if catia:
        cfg.cell(row=r_catia, column=2, value=catia)
    if v2:
        cfg.cell(row=r_v2, column=2, value=v2)
    catia_path = str(cfg.cell(row=r_catia, column=2).value or "").strip().strip('"')
    v2_path = str(cfg.cell(row=r_v2, column=2).value or "").strip().strip('"')

    problems = []
    if not catia_path or not os.path.isfile(catia_path):
        problems.append(f"CONFIG!B{r_catia} (CATIA migration report): "
                        f"'{catia_path}' not found")
    if not v2_path or not os.path.isfile(v2_path):
        problems.append(f"CONFIG!B{r_v2} (SysML v2 source): '{v2_path}' not found")
    if problems:
        raise FileNotFoundError("\n".join(problems))

    elements = read_all_sheet(wb)
    catia_elements, catia_props = parse_catia_report(catia_path)
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with open(v2_path, encoding=enc) as f:
                v2_text = f.read()
            break
        except UnicodeDecodeError:
            continue
    v2m = parse_sysml_v2(v2_text)

    success, missed, stats = classify(elements, catia_elements, catia_props, v2m)

    # ---- sanity check: do the three inputs describe the same model? ------
    roots_csv = {canonical_qn(e["qn"]).split("::")[0]
                 for e in elements if e["qn"]} - {""}
    roots_v2 = {qn.split("::")[0] for qn in v2m.by_qn} - {""}
    warning = ""
    if elements and roots_csv and roots_v2 and not (roots_csv & roots_v2):
        warning = (
            "WARNING: the model exported in ALL_ELEMENTS_1_7 (root packages: "
            + ", ".join(sorted(roots_csv)[:5])
            + ") shares no root package with the SysML v2 source (root packages: "
            + ", ".join(sorted(roots_v2)[:5])
            + "). The CSV, the CATIA report and the v2 source must all come "
            "from the SAME project/scope - these results are meaningless.")
        print("\n*** " + warning + "\n")

    # ---- (re)create the output sheets ------------------------------------
    for name in (SHEET_SUCCESS, SHEET_MISSED, SHEET_SUMMARY):
        if name in wb.sheetnames:
            del wb[name]
    pos = wb.sheetnames.index(SHEET_ALL) + 1

    ws = wb.create_sheet(SHEET_SUCCESS, pos)
    write_table(ws, ["Element MetaType (v1.7)", "Element Name",
                     "v1.7 Qualified Name", "Element ID",
                     "Migrated To (v2 construct)", "v2 Qualified Name",
                     "v2 Line(s)", "Match Type", "Notes"], success)

    ws = wb.create_sheet(SHEET_MISSED, pos + 1)
    write_table(ws, ["Category", "Element MetaType (v1.7)", "Element Name",
                     "v1.7 Qualified Name", "Element ID", "Reason"], missed)
    for row in ws.iter_rows(min_row=2, max_col=1):
        for c in row:
            if c.value in CAT_FILLS:
                c.fill = CAT_FILLS[c.value]

    # ---- MIGRATION_SUMMARY: live Excel formulas over the other sheets -----
    # Counts are COUNTIF/COUNTA formulas so the summary recalculates if rows
    # in MIGRATED_SUCCESS / MISSED_MIGRATIONS are edited during review.
    # (The 'excluded' bucket appears on no sheet, so it is derived by
    # subtraction; the v2-parse and CATIA-report counts have no in-workbook
    # source data and stay static.)
    ws = wb.create_sheet(SHEET_SUMMARY, pos + 2)
    ws.append(["Metric", "Count", "% of v1.7 elements"])
    if warning:
        ws.append([warning, "", ""])
    r0 = ws.max_row + 1          # row of the 'total' line
    T = f"$B${r0}"               # absolute ref to the total count

    def frow(label, count_formula, row):
        ws.append([label, count_formula, f'=IFERROR(B{row}/{T},"")'])

    frow("v1.7 elements + diagrams exported (ALL_ELEMENTS_1_7)",
         f"=COUNTA({SHEET_ALL}!A:A)-1", r0)
    frow("    of which diagrams",
         f'=COUNTIF({SHEET_ALL}!E:E,"Diagram")', r0 + 1)
    frow("MIGRATED_SUCCESS - matched by qualified name",
         f'=COUNTIF({SHEET_SUCCESS}!H:H,"Qualified name")', r0 + 2)
    frow("MIGRATED_SUCCESS - matched by name only",
         f'=COUNTIF({SHEET_SUCCESS}!H:H,"Name*")', r0 + 3)
    frow("    of which also flagged in CATIA report",
         f'=COUNTIF({SHEET_SUCCESS}!I:I,"*CATIA*")', r0 + 4)
    frow("Listed in CATIA not-migrated report (excluded)",
         f"=B{r0}-B{r0 + 2}-B{r0 + 3}-B{r0 + 6}-B{r0 + 7}-B{r0 + 8}", r0 + 5)
    frow("MISSED_MIGRATIONS - MISSING",
         f'=COUNTIF({SHEET_MISSED}!A:A,"MISSING")', r0 + 6)
    frow("MISSED_MIGRATIONS - REFERENCE_ONLY",
         f'=COUNTIF({SHEET_MISSED}!A:A,"REFERENCE_ONLY")', r0 + 7)
    frow("MISSED_MIGRATIONS - UNVERIFIABLE_UNNAMED",
         f'=COUNTIF({SHEET_MISSED}!A:A,"UNVERIFIABLE_UNNAMED")', r0 + 8)
    frow("Accounted for (migrated + CATIA-reported)",
         f"=B{r0 + 2}+B{r0 + 3}+B{r0 + 5}", r0 + 9)
    for extra in (
        ["", "", ""],
        ["SysML v2 named declarations parsed",
         len({r for v in v2m.by_qn.values() for r in v}), ""],
        ["SysML v2 unnamed declarations parsed", sum(v2m.unnamed.values()), ""],
        ["CATIA report: not-migrated elements", len(catia_elements), ""],
        ["CATIA report: elements with lost properties", len(catia_props), ""],
        ["", "", ""],
        ["CATIA report file", catia_path, ""],
        ["SysML v2 source file", v2_path, ""],
        ["Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), ""],
    ):
        ws.append(extra)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT_NAME, size=10)
    for r in range(r0, r0 + 10):
        ws.cell(row=r, column=2).number_format = "0"
        ws.cell(row=r, column=3).number_format = "0.0%"
    style_header(ws, 3)
    autosize(ws, 3)
    ws.column_dimensions["B"].width = 14   # autosize would size B by the
    ws.column_dimensions["C"].width = 18   # formula text, not the result

    wb.save(xlsx_path)
    print(f"Updated workbook: {xlsx_path}")
    print(f"  {SHEET_SUCCESS}:  {len(success)} elements confirmed in the v2 source "
          f"({stats['matched_qn']} by qualified name, {stats['matched_name']} by name)")
    print(f"  {SHEET_MISSED}: {len(missed)} rows "
          f"({stats['missing']} MISSING, {stats['reference_only']} REFERENCE_ONLY, "
          f"{stats['unnamed']} UNVERIFIABLE_UNNAMED)")
    print(f"  Excluded (already in CATIA not-migrated report): {stats['in_catia']}")
    return xlsx_path


# ------------------------------------------- per-metatype summary section

# Fixed metatype list (praxis standard, alphabetical). Metatypes found in
# the model but not listed here are merged in alphabetically at run time;
# the EXCLUDED_METATYPES (defined further down) never appear in the section.
FIXED_METATYPES = [
    "Abstraction", "Accept Event Action", "Activity", "Activity Final Node",
    "Activity Parameter Node", "Activity Partition", "Allocate",
    "Allocated Activity Partition", "Association", "Binding Connector",
    "Block", "Boolean Tagged Value", "Call Behavior Action", "Comment",
    "Connector", "Connector End", "Constraint", "Constraint Block",
    "Constraint Parameter", "Constraint Property", "Control Flow",
    "Decision Node", "Dependency", "DeriveReqt", "Diagram", "Element Import",
    "Element MetaType", "Element Tagged Value", "Final State",
    "Flow Final Node", "Flow Property", "Fork Node", "Functional Requirement",
    "Generalization", "Initial Node", "Input Pin", "Integer Tagged Value",
    "Interface Block", "Item Flow", "Legend", "Legend Item",
    "Literal Integer", "Literal Unlimited Natural", "Merge Node",
    "Object Flow", "Opaque Expression", "Output Pin", "Package",
    "Package Import", "Parameter", "Part Property", "Profile Application",
    "Property", "Proxy Port", "Pseudo State", "Rationale",
    "Reference Property", "Refine", "Region", "Requirement", "Satisfy",
    "Send Signal Action", "Signal", "Signal Event", "State", "State Machine",
    "String Tagged Value", "Test Case", "Trace", "Transition", "Trigger",
    "UseCase", "Value Property", "Value Type", "Verify",
]


def merged_metatypes(found_names):
    """The FIXED list in its given order, with unlisted found metatypes
    inserted at their alphabetical position (case-insensitive)."""
    out = [m for m in FIXED_METATYPES if m.casefold() not in _EXCL_SET]
    known = {m.casefold() for m in out}
    for name in sorted(found_names, key=str.casefold):
        k = name.casefold()
        if k in known:
            continue
        known.add(k)
        pos = len(out)
        for i, m in enumerate(out):
            if m.casefold() > k:
                pos = i
                break
        out.insert(pos, name)
    return out


def _fq(s):
    """Quote a literal string for use inside an Excel formula."""
    return '"' + str(s).replace('"', '""') + '"'


def append_metatype_section(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path)
    for sheet in (SHEET_SUMMARY, SHEET_ALL, SHEET_SUCCESS):
        if sheet not in wb.sheetnames:
            print(f"note: no {sheet} sheet in "
                  f"{os.path.basename(xlsx_path)} - run the comparison "
                  "first; metatype section skipped")
            wb.close()
            return
    ws = wb[SHEET_SUMMARY]
    all_ws = wb[SHEET_ALL]

    # metatypes actually present (first-seen spelling kept for display)
    found = {}
    for (v,) in all_ws.iter_rows(min_row=2, min_col=1, max_col=1,
                                 values_only=True):
        name = str(v).strip() if v is not None else ""
        if name:
            found.setdefault(name.casefold(), name)

    # data-row ranges: never include row 1 headers in the COUNTIFs
    all_last = max(all_ws.max_row, 2)
    suc_last = max(wb[SHEET_SUCCESS].max_row, 2)
    all_rng = f"{SHEET_ALL}!$A$2:$A${all_last}"
    suc_rng = f"{SHEET_SUCCESS}!$A$2:$A${suc_last}"

    r0 = ws.max_row + 2
    for col, title in ((1, "Element Metatype"), (2, "Quantity"),
                       (3, "% Migration Success")):
        c = ws.cell(row=r0, column=col, value=title)
        c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")

    r = r0
    for mt in merged_metatypes(found.values()):
        r += 1
        a = ws.cell(row=r, column=1, value=mt)
        a.font = Font(name=FONT_NAME, size=10)
        b = ws.cell(row=r, column=2,
                    value=f"=COUNTIF({all_rng},{_fq(mt)})")
        b.font = Font(name=FONT_NAME, size=10)
        b.number_format = "0"
        c = ws.cell(row=r, column=3,
                    value=f'=IFERROR(COUNTIF({suc_rng},{_fq(mt)})/B{r},"")')
        c.font = Font(name=FONT_NAME, size=10)
        c.number_format = "0.0%"

    wb.save(xlsx_path)
    print(f"Appended per-metatype section to {SHEET_SUMMARY}: "
          f"{r - r0} metatype rows (header at row {r0}).")


# ----------------------------------------------------- excluded metatypes

EXCLUDED_METATYPES = [
    "Boolean Tagged Value",
    "Diagram",
    "Element Tagged Value",
    "Integer Tagged Value",
    "Literal Unlimited Natural",
    "String Tagged Value",
]
_EXCL_SET = {m.casefold() for m in EXCLUDED_METATYPES}

SHEET_EXCLUDED = "EXCLUDED_METATYPES"


def split_excluded(xlsx_path):
    """Move rows of the excluded metatypes from ALL_ELEMENTS_1_7 to the
    EXCLUDED_METATYPES sheet (idempotent; dedup by Element ID)."""
    wb = openpyxl.load_workbook(xlsx_path)
    if SHEET_ALL not in wb.sheetnames:
        wb.close()
        raise ValueError(f"{os.path.basename(xlsx_path)} has no "
                         f"{SHEET_ALL} sheet")
    ws = wb[SHEET_ALL]
    hm = {str(c.value).strip(): i for i, c in enumerate(ws[1])
          if c.value is not None}

    def norm(row):
        out = []
        for h in ALL_HEADERS:
            j = hm.get(h)
            out.append(row[j] if j is not None and j < len(row) else None)
        return out

    keep, excl = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        r = norm(row)
        mt = str(r[0] or "").strip().casefold()
        (excl if mt in _EXCL_SET else keep).append(r)

    # merge rows excluded on an earlier run (dedup by Element ID, else by
    # the full row tuple for id-less rows)
    if SHEET_EXCLUDED in wb.sheetnames:
        wse = wb[SHEET_EXCLUDED]
        hme = {str(c.value).strip(): i for i, c in enumerate(wse[1])
               if c.value is not None}
        seen_ids = {str(r[3]) for r in excl if r[3]}
        seen_rows = {tuple(str(v) for v in r) for r in excl}
        for row in wse.iter_rows(min_row=2, values_only=True):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            r = [row[hme[h]] if h in hme and hme[h] < len(row) else None
                 for h in ALL_HEADERS]
            rid = str(r[3]) if r[3] else ""
            if (rid and rid in seen_ids) or tuple(str(v) for v in r) in seen_rows:
                continue
            excl.append(r)
        del wb[SHEET_EXCLUDED]

    idx = wb.sheetnames.index(SHEET_ALL)
    del wb[SHEET_ALL]
    ws2 = wb.create_sheet(SHEET_ALL, idx)
    write_table(ws2, ALL_HEADERS, keep)
    wse = wb.create_sheet(SHEET_EXCLUDED, idx + 1)
    write_table(wse, ALL_HEADERS, excl)
    wb.save(xlsx_path)
    wb.close()
    print(f"Excluded metatypes: {len(excl)} element(s) set aside on "
          f"{SHEET_EXCLUDED}; {len(keep)} remain in {SHEET_ALL}.")


def append_summary_note(xlsx_path):
    """One transparent line under the summary: how many elements the
    metrics deliberately ignore."""
    wb = openpyxl.load_workbook(xlsx_path)
    if SHEET_SUMMARY not in wb.sheetnames:
        wb.close()
        return
    ws = wb[SHEET_SUMMARY]
    r = ws.max_row + 1
    a = ws.cell(row=r, column=1,
                value="Excluded from ALL metrics above (see "
                      f"{SHEET_EXCLUDED} sheet): "
                      + ", ".join(EXCLUDED_METATYPES))
    a.font = Font(name=FONT_NAME, size=10, italic=True)
    b = ws.cell(row=r, column=2,
                value=f"=MAX(COUNTA({SHEET_EXCLUDED}!A:A)-1,0)")
    b.font = Font(name=FONT_NAME, size=10)
    b.number_format = "0"
    wb.save(xlsx_path)
    wb.close()


def finish(xlsx_path):
    append_summary_note(xlsx_path)
    append_metatype_section(xlsx_path)


# --------------------------------------------------------------------- main

def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Audit SysML v1.7 -> v2 migration success (CATIA Magic / "
                    "MSOSA 2026x): comparison workbook plus per-metatype "
                    "breakdown, with Tagged Values, Diagrams and Literal "
                    "Unlimited Naturals excluded from all metrics.")
    ap.add_argument("input", nargs="?",
                    help=".csv/.tsv export (build mode) or the "
                         "*_MigrationComparison.xlsx (compare mode)")
    ap.add_argument("--catia", help="path to MigrationReportOf<project>.xlsx")
    ap.add_argument("--v2", help="path to the SysML v2 .txt/.sysml source")
    args = ap.parse_args(argv)

    path = args.input
    if not path:
        path = input("Path to velocity export .csv or comparison .xlsx: "
                     ).strip().strip('"')
    path = path.strip().strip('"')
    if not os.path.isfile(path):
        sys.exit(f"File not found: {path}")

    for label, val in (("--catia", args.catia), ("--v2", args.v2)):
        if val:
            v = val.strip().strip('"')
            if not os.path.isfile(v):
                sys.exit(f"{label} file not found: {v}")

    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv", ".txt") and not path.lower().endswith(".sysml"):
        out = build_workbook(path, catia=args.catia, v2=args.v2)
        split_excluded(out)
        wb = openpyxl.load_workbook(out, read_only=True)
        cfg = wb[SHEET_CONFIG]
        b3 = str(cfg["B3"].value or "").strip().strip('"')
        b4 = str(cfg["B4"].value or "").strip().strip('"')
        wb.close()
        if os.path.isfile(b3) and os.path.isfile(b4):
            print("Both CONFIG paths resolved - running comparison now...")
            run_comparison(out)
            finish(out)
        else:
            print("Comparison did not run (CONFIG paths incomplete). Fill "
                  "CONFIG!B3 / CONFIG!B4, then rerun this script on the "
                  "workbook.")
    elif ext in (".xlsx", ".xlsm"):
        split_excluded(path)
        out = run_comparison(path, catia=args.catia, v2=args.v2)
        finish(out or path)
    else:
        sys.exit("Input must be a .csv/.tsv (build mode) or .xlsx (compare mode)")


if __name__ == "__main__":
    main()
